import sqlite3
from datetime import datetime, timedelta
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
import subprocess
import os
import sys
import shutil
import threading
import struct
import tempfile

# --- ROTINA DE IMPRESSÃO E EMISSÃO DE COMPROVANTES (ADEGA GONE DRACK) ---
def imprimir_comprovante_venda(vendedor, carrinho, total, pago, troco, pagamentos):
    """Gera um arquivo de texto formatado como cupom não fiscal e envia para a impressora padrão."""
    data_hora = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    
    cupom = []
    cupom.append("=" * 33)
    cupom.append("       ADEGA GONE DRACK       ")
    cupom.append("   Comprovante de Venda (PDV)  ")
    cupom.append("=" * 33)
    cupom.append(f"Data/Hora: {data_hora}")
    cupom.append(f"Atendente: {vendedor}")
    cupom.append("-" * 33)
    cupom.append(f"{'ITEM':<18} {'QTD':>4} {'TOTAL':>9}")
    cupom.append("-" * 33)

    for item in carrinho:
        nome_prod = item['nome'][:18]
        qtd = f"{item['qtd']}x"
        subtotal = f"R${item['sub']:.2f}"
        cupom.append(f"{nome_prod:<18} {qtd:>4} {subtotal:>9}")

    cupom.append("-" * 33)
    cupom.append(f"TOTAL:             R$ {total:>8.2f}")
    
    for p in pagamentos:
        forma = p['forma'].upper()
        val = p['valor']
        cupom.append(f"PAGO ({forma}):    R$ {val:>8.2f}")
        
    if troco > 0:
        cupom.append(f"TROCO:             R$ {troco:>8.2f}")
        
    cupom.append("=" * 33)
    cupom.append("    Obrigado pela preferencia!   ")
    cupom.append("=" * 33)
    cupom.append("\n\n")

    texto_cupom = "\n".join(cupom)

    try:
        with tempfile.NamedTemporaryFile(mode="w", delete=False, suffix=".txt", encoding="utf-8") as temp_file:
            temp_file.write(texto_cupom)
            caminho_temp = temp_file.name

        if sys.platform.startswith('linux'):
            subprocess.run(['lpr', caminho_temp])
        elif sys.platform.startswith('win'):
            os.startfile(caminho_temp, "print")
    except Exception as e:
        print(f"[ERRO IMPRESSAO] Nao foi possivel imprimir: {e}")

def gerar_pdf_comprovante_venda(caminho_pdf, vendedor, carrinho, total, pago, troco, pagamentos):
    """Gera um arquivo PDF com o comprovante de venda."""
    data_hora = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    c = canvas.Canvas(caminho_pdf, pagesize=letter)
    y = 750
    
    c.setFont("Courier-Bold", 12)
    c.drawString(50, y, "=================================") ; y -= 15
    c.drawString(50, y, "       ADEGA GONE DRACK          ") ; y -= 15
    c.drawString(50, y, "   Comprovante de Venda (PDV)    ") ; y -= 15
    c.drawString(50, y, "=================================") ; y -= 20
    
    c.setFont("Courier", 10)
    c.drawString(50, y, f"Data/Hora: {data_hora}") ; y -= 15
    c.drawString(50, y, f"Atendente: {vendedor}") ; y -= 15
    c.drawString(50, y, "---------------------------------") ; y -= 15
    c.drawString(50, y, f"{'ITEM':<18} {'QTD':>4} {'TOTAL':>9}") ; y -= 15
    c.drawString(50, y, "---------------------------------") ; y -= 15

    for item in carrinho:
        if y < 50:
            c.showPage()
            y = 750
            c.setFont("Courier", 10)
        nome_prod = item['nome'][:18]
        qtd = f"{item['qtd']}x"
        subtotal = f"R${item['sub']:.2f}"
        c.drawString(50, y, f"{nome_prod:<18} {qtd:>4} {subtotal:>9}")
        y -= 15

    c.drawString(50, y, "---------------------------------") ; y -= 15
    c.setFont("Courier-Bold", 10)
    c.drawString(50, y, f"TOTAL:             R$ {total:>8.2f}") ; y -= 15
    
    c.setFont("Courier", 10)
    for p in pagamentos:
        forma = p['forma'].upper()
        val = p['valor']
        c.drawString(50, y, f"PAGO ({forma}):    R$ {val:>8.2f}")
        y -= 15
        
    if troco > 0:
        c.drawString(50, y, f"TROCO:             R$ {troco:>8.2f}") ; y -= 15
        
    c.drawString(50, y, "=================================") ; y -= 15
    c.drawString(50, y, "    Obrigado pela preferencia!   ") ; y -= 15
    c.drawString(50, y, "=================================")
    
    c.save()

# --- ROTINA DE BACKUP AUTOMÁTICO (GOOGLE DRIVE / RCLONE) ---
def fazer_backup_nuvem():
    """Realiza a cópia do banco de dados e sincroniza com o Google Drive."""
    def copiar():
        try:
            home_dir = os.path.expanduser("~")

            if sys.platform.startswith('linux'):
                pasta_destino = os.path.join(home_dir, "GoogleDrive", "AdegaBackup")
                executavel_rclone = "rclone"
            elif sys.platform.startswith('win'):
                pasta_destino = os.path.join(home_dir, "Documents", "AdegaBackup")
                is_64bit = struct.calcsize("P") * 8 == 64
                nome_exe = "rclone64.exe" if is_64bit else "rclone32.exe"
                caminho_local_rclone = os.path.join(os.getcwd(), nome_exe)
                if os.path.exists(caminho_local_rclone):
                    executavel_rclone = caminho_local_rclone
                else:
                    executavel_rclone = nome_exe
            else:
                pasta_destino = os.path.join(home_dir, "AdegaBackup")
                executavel_rclone = "rclone"

            if not os.path.exists(pasta_destino):
                os.makedirs(pasta_destino)

            shutil.copy2("adega.db", os.path.join(pasta_destino, "adega_backup.db"))
            data_str = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
            shutil.copy2("adega.db", os.path.join(pasta_destino, f"adega_{data_str}.db"))

            comando_rclone = [
                executavel_rclone, "sync", pasta_destino, "gdrive:AdegaBackup",
                "--tpslimit", "5",
                "--fast-list"
            ]
            
            resultado = subprocess.run(comando_rclone, capture_output=True, text=True)

            if resultado.returncode == 0:
                print(f"[BACKUP NUVEM OK] Sincronizado com o Google Drive às {data_str}")
            else:
                print(f"[AVISO NUVEM] Falha pontual na sincronização remota. Backup local preservado.")

        except Exception as e:
            print(f"[ERRO BACKUP] Falha ao realizar backup: {e}")

    threading.Thread(target=copiar, daemon=True).start()

# --- FUNÇÕES DE EXPORTAÇÃO DE RELATÓRIO ---
def exportar_relatorio_docx(caminho_arquivo, titulo, periodo, faturamento, lucro, metodos, texto_detalhes):
    doc = Document()
    doc.add_heading(f'Adega Gone Drack - {titulo}', level=1)
    if periodo:
        doc.add_paragraph(f"Período: {periodo}")
    if faturamento is not None:
        doc.add_paragraph(f"Faturamento Total: R$ {faturamento:.2f}")
    if lucro is not None:
        doc.add_paragraph(f"Lucro Líquido: R$ {lucro:.2f}")
    
    if metodos:
        doc.add_heading('Resumo por Método de Pagamento', level=2)
        for metodo, valor in metodos.items():
            doc.add_paragraph(f"{metodo}: R$ {valor:.2f}")
        
    doc.add_heading('Detalhamento', level=2)
    doc.add_paragraph(texto_detalhes)
    
    doc.save(caminho_arquivo)

def exportar_relatorio_pdf(caminho_arquivo, titulo, periodo, faturamento, lucro, metodos, texto_detalhes):
    c = canvas.Canvas(caminho_arquivo, pagesize=letter)
    y = 750
    
    c.setFont("Helvetica-Bold", 16)
    c.drawString(50, y, f"Adega Gone Drack - {titulo}")
    y -= 25
    
    c.setFont("Helvetica", 12)
    if periodo:
        c.drawString(50, y, f"Período: {periodo}")
        y -= 20
    if faturamento is not None and lucro is not None:
        c.drawString(50, y, f"Faturamento Total: R$ {faturamento:.2f} | Lucro: R$ {lucro:.2f}")
        y -= 30
    
    if metodos:
        c.setFont("Helvetica-Bold", 13)
        c.drawString(50, y, "Métodos de Pagamento:")
        y -= 20
        c.setFont("Helvetica", 11)
        for metodo, valor in metodos.items():
            c.drawString(70, y, f"{metodo}: R$ {valor:.2f}")
            y -= 15
        y -= 20
        
    c.setFont("Helvetica-Bold", 13)
    c.drawString(50, y, "Detalhamento:")
    y -= 20
    
    c.setFont("Courier", 10)
    for linha in texto_detalhes.split("\n"):
        if y < 50:
            c.showPage()
            y = 750
            c.setFont("Courier", 10)
        c.drawString(50, y, linha)
        y -= 15
        
    c.save()

def abrir_arquivo_gerado(caminho):
    try:
        if sys.platform.startswith('linux'):
            subprocess.run(['xdg-open', caminho])
        elif sys.platform.startswith('win'):
            os.startfile(caminho)
        elif sys.platform.startswith('darwin'):
            subprocess.run(['open', caminho])
    except Exception as e:
        print(f"[LOG ERRO] Erro ao abrir arquivo: {e}")

# --- CONEXÃO E ESTRUTURA DO BANCO ---
def conectar():
    conexao = sqlite3.connect('adega.db')
    cursor = conexao.cursor()
    
    cursor.execute('''CREATE TABLE IF NOT EXISTS produtos 
        (id INTEGER PRIMARY KEY AUTOINCREMENT, nome TEXT, quantidade INTEGER, 
        preco_custo REAL, preco_venda REAL, preco_fardo REAL, unidades_por_fardo INTEGER, 
        codigo_barras TEXT, categoria TEXT)''')
    
    cursor.execute('''CREATE TABLE IF NOT EXISTS vendas 
        (id INTEGER PRIMARY KEY AUTOINCREMENT, id_produto INTEGER, quantidade_vendida INTEGER, 
        valor_pago REAL, tipo_venda TEXT, metodo_pagamento TEXT, custo_na_venda REAL, data_venda TEXT, vendedor TEXT)''')
    
    cursor.execute('''CREATE TABLE IF NOT EXISTS pagamentos_venda 
        (id INTEGER PRIMARY KEY AUTOINCREMENT, forma_pagamento TEXT, valor REAL, usuario TEXT, data_pagamento TEXT)''')

    cursor.execute('''CREATE TABLE IF NOT EXISTS config (chave TEXT PRIMARY KEY, valor TEXT)''')
    cursor.execute('''CREATE TABLE IF NOT EXISTS usuarios 
        (id INTEGER PRIMARY KEY AUTOINCREMENT, nome TEXT, login TEXT UNIQUE, senha TEXT, nivel TEXT)''')
    cursor.execute('''CREATE TABLE IF NOT EXISTS logs 
        (id INTEGER PRIMARY KEY AUTOINCREMENT, usuario_nome TEXT, acao TEXT, data_hora TEXT)''')

    cursor.execute('''CREATE TABLE IF NOT EXISTS vendas_avulsas 
        (id INTEGER PRIMARY KEY AUTOINCREMENT, nome_informado TEXT, quantidade INTEGER, 
        valor_pago REAL, data_venda TEXT, vendedor TEXT)''')

    cursor.execute("PRAGMA table_info(vendas)")
    if 'vendedor' not in [col[1] for col in cursor.fetchall()]:
        cursor.execute("ALTER TABLE vendas ADD COLUMN vendedor TEXT DEFAULT 'Sistema'")
    
    cursor.execute("SELECT COUNT(*) FROM usuarios")
    if cursor.fetchone()[0] == 0:
        cursor.execute("INSERT INTO usuarios (nome, login, senha, nivel) VALUES (?, ?, ?, ?)",
                       ('Administrador', 'admin', '1234', 'Admin'))
    
    conexao.commit()
    return conexao

def verificar_login(usuario, senha):
    conn = conectar(); cursor = conn.cursor()
    cursor.execute("SELECT nome, nivel FROM usuarios WHERE login = ? AND senha = ?", (usuario, senha))
    res = cursor.fetchone(); conn.close()
    return {"nome": res[0], "nivel": res[1]} if res else None

def registrar_venda(id_p, qtd, tipo, metodo, vendedor="Sistema", nome_avulso=None, preco_avulso=None):
    conn = conectar(); cursor = conn.cursor()
    
    if id_p == 0 and nome_avulso:
        val = preco_avulso * qtd
        c_tot = 0.0
        cursor.execute("INSERT INTO vendas (id_produto, quantidade_vendida, tipo_venda, metodo_pagamento, custo_na_venda, valor_pago, data_venda, vendedor) VALUES (?,?,?,?,?,?, datetime('now','localtime'), ?)",
                       (0, qtd, tipo, metodo, c_tot, val, vendedor))
        cursor.execute("INSERT INTO vendas_avulsas (nome_informado, quantidade, valor_pago, data_venda, vendedor) VALUES (?,?,?, datetime('now','localtime'), ?)",
                       (nome_avulso, qtd, val, vendedor))
    else:
        cursor.execute("SELECT preco_custo, preco_venda, preco_fardo, unidades_por_fardo FROM produtos WHERE id = ?", (id_p,))
        p = cursor.fetchone()
        if not p: return
        un = p[3] if p[3] else 1
        if tipo in ['Fardo', 'Caixa', 'Pacote']:
            val, q_est, c_tot = (p[2]*qtd), (qtd*un), ((p[0]*un)*qtd)
        else:
            val, q_est, c_tot = (p[1]*qtd), qtd, (p[0]*qtd)
        cursor.execute("INSERT INTO vendas (id_produto, quantidade_vendida, tipo_venda, metodo_pagamento, custo_na_venda, valor_pago, data_venda, vendedor) VALUES (?,?,?,?,?,?, datetime('now','localtime'), ?)",
                       (id_p, qtd, tipo, metodo, c_tot, val, vendedor))
        cursor.execute("UPDATE produtos SET quantidade = quantidade - ? WHERE id = ?", (q_est, id_p))
        
    conn.commit(); conn.close()
    fazer_backup_nuvem()

def registrar_pagamento_detalhado(forma, valor, usuario="Sistema"):
    conn = conectar(); cursor = conn.cursor()
    cursor.execute("INSERT INTO pagamentos_venda (forma_pagamento, valor, usuario, data_pagamento) VALUES (?, ?, ?, datetime('now','localtime'))",
                   (forma, valor, usuario))
    conn.commit(); conn.close()

def calcular_lucro_hoje(d1, d2):
    conn = conectar(); cursor = conn.cursor()
    cursor.execute('SELECT SUM(valor_pago - custo_na_venda) FROM vendas WHERE data_venda BETWEEN ? AND ?', (d1, d2))
    res = cursor.fetchone()[0]; conn.close(); return res if res else 0.0

def resumo_vendas_por_metodo(met, d1, d2):
    conn = conectar(); cursor = conn.cursor()
    cursor.execute("SELECT SUM(valor_pago) FROM vendas WHERE metodo_pagamento = ? AND data_venda BETWEEN ? AND ?", (met, d1, d2))
    res = cursor.fetchone()[0]; conn.close(); return res if res else 0.0

def resumo_vendas_por_vendedor(d1, d2):
    conn = conectar(); cursor = conn.cursor()
    cursor.execute('SELECT vendedor, SUM(valor_pago) FROM vendas WHERE data_venda BETWEEN ? AND ? GROUP BY vendedor', (d1, d2))
    res = cursor.fetchall(); conn.close(); return res

def produtos_mais_vendidos_hoje(d1, d2):
    conn = conectar(); cursor = conn.cursor()
    cursor.execute('''
        SELECT COALESCE(p.nome, '[AVULSO] Item Sem Cadastro'), SUM(v.quantidade_vendida), v.tipo_venda 
        FROM vendas v 
        LEFT JOIN produtos p ON v.id_produto = p.id 
        WHERE v.data_venda BETWEEN ? AND ? 
        GROUP BY COALESCE(p.nome, '[AVULSO] Item Sem Cadastro'), v.tipo_venda
    ''', (d1, d2))
    res = cursor.fetchall(); conn.close(); return res

def vendas_por_filtro_produto(nome_produto, d1, d2):
    conn = conectar(); cursor = conn.cursor()
    termo = f"%{nome_produto.strip().lower()}%"
    cursor.execute('''
        SELECT COALESCE(p.nome, '[AVULSO] Item Sem Cadastro'), SUM(v.quantidade_vendida), v.tipo_venda, SUM(v.valor_pago)
        FROM vendas v 
        LEFT JOIN produtos p ON v.id_produto = p.id 
        WHERE LOWER(COALESCE(p.nome, '[AVULSO] Item Sem Cadastro')) LIKE ? AND v.data_venda BETWEEN ? AND ?
        GROUP BY COALESCE(p.nome, '[AVULSO] Item Sem Cadastro'), v.tipo_venda
    ''', (termo, d1, d2))
    res = cursor.fetchall(); conn.close()
    return res

def listar_produtos_sem_cadastro(d1, d2):
    conn = conectar(); cursor = conn.cursor()
    cursor.execute('''
        SELECT nome_informado, SUM(quantidade), SUM(valor_pago), vendedor
        FROM vendas_avulsas
        WHERE data_venda BETWEEN ? AND ?
        GROUP BY nome_informado, vendedor
        ORDER BY SUM(quantidade) DESC
    ''', (d1, d2))
    res = cursor.fetchall(); conn.close(); return res

def listar_produtos_reposicao(limite=10):
    conn = conectar(); cursor = conn.cursor()
    cursor.execute('SELECT id, nome, categoria, quantidade, unidades_por_fardo FROM produtos WHERE quantidade <= ? ORDER BY quantidade ASC', (limite,))
    res = cursor.fetchall(); conn.close(); return res

def obter_horarios():
    conn = conectar(); cursor = conn.cursor()
    cursor.execute("SELECT valor FROM config WHERE chave = 'hora_abertura'"); ab = cursor.fetchone()
    cursor.execute("SELECT valor FROM config WHERE chave = 'hora_fechamento'"); fe = cursor.fetchone()
    conn.close(); return (ab[0] if ab else "08:00", fe[0] if fe else "22:00")

def configurar_horarios(ab, fe):
    conn = conectar(); cursor = conn.cursor()
    cursor.execute("INSERT OR REPLACE INTO config (chave, valor) VALUES (?, ?)", ("hora_abertura", ab))
    cursor.execute("INSERT OR REPLACE INTO config (chave, valor) VALUES (?, ?)", ("hora_fechamento", fe))
    conn.commit(); conn.close()

def adicionar_usuario(nome, login, senha, nivel):
    conn = conectar(); cursor = conn.cursor()
    cursor.execute("INSERT INTO usuarios (nome, login, senha, nivel) VALUES (?, ?, ?, ?)", (nome, login, senha, nivel))
    conn.commit(); conn.close()

def listar_usuarios():
    conn = conectar(); cursor = conn.cursor(); cursor.execute("SELECT id, nome, login, nivel FROM usuarios"); res = cursor.fetchall(); conn.close(); return res

def excluir_usuario_db(id_u):
    conn = conectar(); cursor = conn.cursor(); cursor.execute("DELETE FROM usuarios WHERE id = ?", (id_u,)); conn.commit(); conn.close()

def adicionar_produto(nome, quantidade, custo, venda_unit, venda_fardo=None, qtd_fardo=None, id_produto=None, codigo_barras=None, categoria=None):
    conexao = conectar(); cursor = conexao.cursor()
    if id_produto:
        cursor.execute('''UPDATE produtos SET nome=?, quantidade=?, preco_custo=?, preco_venda=?, preco_fardo=?, unidades_por_fardo=?, codigo_barras=?, categoria=? WHERE id=?''', (nome, quantidade, custo, venda_unit, venda_fardo, qtd_fardo, codigo_barras, categoria, id_produto))
    else:
        cursor.execute('''INSERT INTO produtos (nome, quantidade, preco_custo, preco_venda, preco_fardo, unidades_por_fardo, codigo_barras, categoria) VALUES (?,?,?,?,?,?,?,?)''', (nome, quantidade, custo, venda_unit, venda_fardo, qtd_fardo, codigo_barras, categoria))
    conexao.commit(); conexao.close()

def buscar_produto_por_id(id_p):
    conn = conectar(); cursor = conn.cursor(); cursor.execute("SELECT * FROM produtos WHERE id = ?", (id_p,)); res = cursor.fetchone(); conn.close(); return res

def buscar_produto_por_codigo(codigo):
    conn = conectar(); cursor = conn.cursor(); cursor.execute("SELECT * FROM produtos WHERE codigo_barras = ?", (codigo,)); res = cursor.fetchone(); conn.close(); return res

def buscar_produto_por_nome(nome):
    conn = conectar(); cursor = conn.cursor(); cursor.execute("SELECT id, nome, preco_venda, preco_fardo, quantidade FROM produtos WHERE nome LIKE ?", ('%'+nome+'%',)); res = cursor.fetchall(); conn.close(); return res

def buscar_id_por_nome_exato(nome):
    conn = conectar(); cursor = conn.cursor(); cursor.execute("SELECT id FROM produtos WHERE nome = ?", (nome,)); res = cursor.fetchone(); conn.close(); return res[0] if res else None

def listar_produtos():
    conn = conectar(); cursor = conn.cursor(); cursor.execute('SELECT * FROM produtos'); res = cursor.fetchall(); conn.close(); return res

def excluir_produto_db(id_p):
    conn = conectar(); cursor = conn.cursor(); cursor.execute("DELETE FROM produtos WHERE id = ?", (id_p,)); conn.commit(); conn.close()

def listar_categorias_unicas():
    conn = conectar(); cursor = conn.cursor(); cursor.execute("SELECT DISTINCT categoria FROM produtos WHERE categoria IS NOT NULL AND categoria != '' ORDER BY categoria ASC"); res = cursor.fetchall(); conn.close(); return [item[0] for item in res]